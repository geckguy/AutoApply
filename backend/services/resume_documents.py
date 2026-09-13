"""Generate private, downloadable PDF and DOCX resume versions."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

from backend.models.profile import UserProfile
from backend.services.llm_client import LLMResponseError


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-.")
    return cleaned[:100] or "resume"


def _lines(value: str | None) -> list[str]:
    if not value:
        return []
    parts = re.split(r"(?:\r?\n|\s*[•·]\s*)", value)
    return [part.strip(" -\t") for part in parts if part.strip(" -\t")]


def apply_tailoring(
    profile: UserProfile, tailoring: dict[str, Any] | None
) -> UserProfile:
    """Apply AI reordering/rephrasing to a copy while preserving profile facts.

    Raises LLMResponseError when the provider did not return a tailoring object,
    so the caller reports a failed tailoring instead of storing an untailored
    resume version as a success.
    """
    if not isinstance(tailoring, dict):
        raise LLMResponseError("Tailoring provider returned an invalid document")
    result = profile.model_copy(deep=True)
    summary = tailoring.get("summary")
    if isinstance(summary, str) and summary.strip():
        result.personal.summary = summary.strip()

    highlighted = [str(skill).strip() for skill in tailoring.get("highlighted_skills", []) if str(skill).strip()]
    known = {skill.casefold(): skill for skill in result.skills}
    ordered = [known[skill.casefold()] for skill in highlighted if skill.casefold() in known]
    result.skills = ordered + [skill for skill in result.skills if skill.casefold() not in {item.casefold() for item in ordered}]

    for suggestion in tailoring.get("experience_bullets", []):
        if not isinstance(suggestion, dict):
            continue
        company = str(suggestion.get("company") or "").strip().casefold()
        title = str(suggestion.get("title") or "").strip().casefold()
        bullets = [str(item).strip() for item in suggestion.get("bullets", []) if str(item).strip()]
        if not bullets:
            continue
        # Require a matching role; never create employment history from model output.
        match = next(
            (
                role for role in result.work_experience
                if company == role.company.strip().casefold()
                and (not title or title == role.title.strip().casefold())
            ),
            None,
        )
        if match:
            match.description = "\n".join(bullets)
    return result


class ResumeDocumentGenerator:
    """Render one structured profile to matching PDF and DOCX files."""

    def __init__(self, root: Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.root.chmod(0o700)

    def render(self, profile: UserProfile, version_id: str, display_name: str) -> dict[str, str]:
        version_dir = self.root / _safe_name(version_id)
        version_dir.mkdir(parents=True, exist_ok=True)
        version_dir.chmod(0o700)
        stem = _safe_name(display_name)
        pdf_path = version_dir / f"{stem}.pdf"
        docx_path = version_dir / f"{stem}.docx"
        self._render_pdf(profile, pdf_path)
        self._render_docx(profile, docx_path)
        pdf_path.chmod(0o600)
        docx_path.chmod(0o600)
        return {"pdf_path": str(pdf_path), "docx_path": str(docx_path)}

    @staticmethod
    def _contact_line(profile: UserProfile) -> str:
        p = profile.personal
        location = ", ".join(filter(None, (p.address.city, p.address.state, p.address.country)))
        return "  |  ".join(filter(None, (p.email, p.phone, location, p.linkedin, p.github, p.portfolio)))

    @classmethod
    def _render_docx(cls, profile: UserProfile, path: Path) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(9.5)

        p = profile.personal
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(" ".join(filter(None, (p.first_name, p.last_name))) or "Resume")
        run.bold = True
        run.font.size = Pt(20)
        contact = document.add_paragraph(cls._contact_line(profile))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def section_heading(text: str) -> None:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(10.5)

        if p.summary:
            section_heading("Summary")
            document.add_paragraph(p.summary)
        if profile.skills:
            section_heading("Skills")
            document.add_paragraph(" • ".join(profile.skills))
        if profile.work_experience:
            section_heading("Experience")
            for role in profile.work_experience:
                dates = " – ".join(filter(None, (role.start_date, role.end_date)))
                title = document.add_paragraph()
                title.paragraph_format.space_after = Pt(0)
                left = title.add_run(f"{role.title} — {role.company}")
                left.bold = True
                if dates:
                    title.add_run(f"   {dates}").italic = True
                for bullet in _lines(role.description):
                    document.add_paragraph(bullet, style="List Bullet")
        if profile.education:
            section_heading("Education")
            for item in profile.education:
                degree = ", ".join(filter(None, (item.degree, item.field)))
                dates = " – ".join(filter(None, (item.start_date, item.end_date)))
                document.add_paragraph(" | ".join(filter(None, (item.institution, degree, dates))))
        if profile.projects:
            section_heading("Projects")
            for project in profile.projects:
                paragraph = document.add_paragraph()
                paragraph.add_run(project.name).bold = True
                if project.description:
                    paragraph.add_run(f" — {project.description}")
        if profile.certifications:
            section_heading("Certifications")
            document.add_paragraph(" • ".join(profile.certifications))
        document.save(path)

    @classmethod
    def _render_pdf(cls, profile: UserProfile, path: Path) -> None:
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "ResumeBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.8,
            leading=11, spaceAfter=3, textColor=colors.HexColor("#18212f"),
        )
        section = ParagraphStyle(
            "ResumeSection", parent=body, fontName="Helvetica-Bold", fontSize=10,
            leading=12, spaceBefore=6, spaceAfter=3, textColor=colors.HexColor("#21506b"),
        )
        name = ParagraphStyle(
            "ResumeName", parent=body, fontName="Helvetica-Bold", fontSize=18,
            leading=21, alignment=TA_CENTER, spaceAfter=2,
        )
        contact = ParagraphStyle("ResumeContact", parent=body, fontSize=8, alignment=TA_CENTER, spaceAfter=7)
        document = SimpleDocTemplate(
            str(path), pagesize=LETTER,
            leftMargin=0.6 * inch, rightMargin=0.6 * inch,
            topMargin=0.48 * inch, bottomMargin=0.48 * inch,
            title="Resume",
        )
        story = []
        p = profile.personal
        full_name = " ".join(filter(None, (p.first_name, p.last_name))) or "Resume"
        story.append(Paragraph(escape(full_name), name))
        story.append(Paragraph(escape(cls._contact_line(profile)), contact))

        def heading(value: str) -> None:
            story.append(Paragraph(escape(value.upper()), section))

        if p.summary:
            heading("Summary")
            story.append(Paragraph(escape(p.summary), body))
        if profile.skills:
            heading("Skills")
            story.append(Paragraph(escape(" • ".join(profile.skills)), body))
        if profile.work_experience:
            heading("Experience")
            for role in profile.work_experience:
                dates = " – ".join(filter(None, (role.start_date, role.end_date)))
                line = f"<b>{escape(role.title)} — {escape(role.company)}</b>"
                if dates:
                    line += f" &nbsp;&nbsp; <i>{escape(dates)}</i>"
                story.append(Paragraph(line, body))
                for bullet in _lines(role.description):
                    story.append(Paragraph(f"• &nbsp;{escape(bullet)}", body))
        if profile.education:
            heading("Education")
            for item in profile.education:
                degree = ", ".join(filter(None, (item.degree, item.field)))
                dates = " – ".join(filter(None, (item.start_date, item.end_date)))
                story.append(Paragraph(escape(" | ".join(filter(None, (item.institution, degree, dates)))), body))
        if profile.projects:
            heading("Projects")
            for project in profile.projects:
                detail = f"<b>{escape(project.name)}</b>"
                if project.description:
                    detail += f" — {escape(project.description)}"
                story.append(Paragraph(detail, body))
        if profile.certifications:
            heading("Certifications")
            story.append(Paragraph(escape(" • ".join(profile.certifications)), body))
        story.append(Spacer(1, 1))
        document.build(story)
